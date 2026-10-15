import copy
import json
import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import backend.app as api
from backend.extract import FIELDS
from backend.locate import locate
from backend.validate import amount, validate


def make_doc(items=(("Room rent", "2", "1500", "3000"), ("CBC test", "1", "450", "450")),
             **values):
    base = {"patient_name": "Ravi Kumar", "hospital_name": "City Care Hospital",
            "bill_date": "10-10-2026", "subtotal": "3450", "tax": "0", "total": "3450"}
    base.update(values)
    return {
        "document_type": "bill",
        "fields": {f: {"value": base.get(f, ""), "confidence": 0.95} for f in FIELDS},
        "line_items": [{"description": d, "qty": q, "rate": r, "amount": a, "confidence": 0.95}
                       for d, q, r, a in items],
    }


def rules(doc):
    return sorted(i["rule"] for i in validate(doc))


def test_amount_parsing():
    assert amount("₹1,23,456.50") == 123456.5
    assert amount("Rs. 450/-") == 450
    assert amount("") is None and amount(None) is None


def test_clean_bill_passes():
    assert validate(make_doc()) == []


def test_wrong_total_is_caught():
    issues = validate(make_doc(total="5200"))
    assert [i["rule"] for i in issues] == ["total_math"]
    assert "5,200.00" in issues[0]["message"]


def test_items_not_matching_subtotal():
    assert rules(make_doc(subtotal="4000", total="4000")) == ["items_sum"]


def test_discount_and_tax():
    assert validate(make_doc(discount="450", tax="150", total="3150")) == []


def test_line_math_missing_field_duplicate_and_low_confidence():
    doc = make_doc(items=(("Paracetamol", "10", "2", "25"), ("Paracetamol", "10", "2", "25")),
                   patient_name="", subtotal="50", total="50")
    doc["fields"]["hospital_name"]["confidence"] = 0.4
    assert rules(doc) == ["duplicate", "line_math", "line_math", "low_confidence", "required"]


def test_rounding_tolerance():
    assert validate(make_doc(total="3450.60")) == []


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "DB", tmp_path / "test.db")
    monkeypatch.setattr(api, "UPLOADS", tmp_path / "uploads")
    monkeypatch.setattr(api, "extract", lambda data, mime: copy.deepcopy(make_doc(total="5200")))
    return TestClient(api.app, headers={"X-Hayara": "1"})


def test_upload_review_correct_export(client):
    r = client.post("/upload", files={"file": ("bill.jpg", b"fake-image", "image/jpeg")})
    assert r.status_code == 200 and r.json()["status"] == "review"
    bill_id = r.json()["id"]

    assert client.get("/bills/1/file").content == b"fake-image"
    assert client.get("/bills", params={"status": "review"}).json()[0]["issue_count"] == 1

    r = client.patch(f"/bills/{bill_id}", json={"field": "total", "value": "3450"})
    assert r.json()["status"] == "approved"
    assert client.patch(f"/bills/{bill_id}", json={"field": "../etc", "value": "x"}).status_code == 422

    with api.db() as conn:
        fix = conn.execute("SELECT ai_value, corrected_value FROM corrections").fetchone()
    assert tuple(fix) == ("5200", "3450")

    assert client.get("/export", params={"format": "csv"}).text.splitlines()[1].startswith("1,bill.jpg,approved")
    assert client.get("/export").content[:2] == b"PK"  # xlsx is a zip
    assert client.get("/export", params={"format": "json"}).json()[0]["data"]["fields"]["total"]["value"] == "3450"


def test_rejects_bad_type(client):
    assert client.post("/upload", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 415


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract not installed")
def test_click_to_source_on_demo_bill():
    root = Path(__file__).resolve().parent.parent / "data"
    doc = json.loads((root / "truth" / "bill_03.json").read_text())
    boxes = locate((root / "samples" / "bill_03.png").read_bytes(), doc)
    assert boxes["total"][0] > 0.7 and boxes["total"][1] > boxes["subtotal"][1]  # bottom right, below subtotal
    assert boxes["line_items.3"][1] > boxes["line_items.1"][1]  # the repeated medicine gets its own row


def test_eval_compare_counts_mistakes():
    from eval import compare
    truth = json.loads((Path(__file__).resolve().parent.parent / "data" / "truth" / "bill_02.json").read_text())
    pred = copy.deepcopy(truth)
    assert compare(pred, truth)["fields_ok"] == compare(pred, truth)["fields"]
    assert compare(pred, truth)["verdict_ok"]  # the planted wrong total is flagged

    pred["fields"]["bill_date"]["value"] = "11/10/2026"     # same date, other format: still right
    pred["fields"]["total"]["value"] = "₹5,730.00"         # misread total: wrong, and hides the mistake
    pred["line_items"][0]["description"] = "Medicines"     # misread row: counted wrong, no math error
    r = compare(pred, truth)
    assert r["wrong"] == ["total"] and r["items_ok"] == r["items"] - 1 and not r["verdict_ok"]


def test_password_protects_hosted_app(client, monkeypatch):
    import base64
    monkeypatch.setenv("APP_PASSWORD", "judges2026")
    assert client.get("/bills").status_code == 401
    assert client.get("/").status_code == 401
    bad = base64.b64encode(b"x:wrong").decode()
    assert client.get("/bills", headers={"Authorization": f"Basic {bad}"}).status_code == 401
    assert client.get("/bills", headers={"Authorization": "Basic !!notbase64"}).status_code == 401
    good = base64.b64encode(b"anyone:judges2026").decode()
    assert client.get("/bills", headers={"Authorization": f"Basic {good}"}).status_code == 200


def test_blank_rows_are_not_duplicates_and_do_not_crash():
    doc = make_doc(items=(("Room rent", "2", "1500", "3000"), ("CBC test", "1", "450", "450"),
                          ("Consumables", "", "", ""), ("Consumables", "", "", "")))
    assert validate(doc) == []


def test_bill_with_no_line_items_is_not_auto_approved():
    assert rules(make_doc(items=())) == ["items_sum"]


def test_approval_survives_unrelated_edit_but_not_new_problem(client):
    bill_id = client.post("/upload", files={"file": ("bill.jpg", b"img", "image/jpeg")}).json()["id"]
    assert client.post(f"/bills/{bill_id}/approve").json()["status"] == "approved"   # wrong total, approved as is
    assert client.patch(f"/bills/{bill_id}", json={"field": "doctor", "value": "Dr. Rao"}).json()["status"] == "approved"
    r = client.patch(f"/bills/{bill_id}", json={"field": "subtotal", "value": "9999"}).json()
    assert r["status"] == "review"  # this edit created a new problem


def test_export_keeps_scan_text_as_text(client):
    bill_id = client.post("/upload", files={"file": ("=1+2.png", b"img", "image/png")}).json()["id"]
    client.patch(f"/bills/{bill_id}", json={"field": "doctor", "value": "=HYPERLINK(\"x\")\x0b"})
    from openpyxl import load_workbook
    import io
    sheet = load_workbook(io.BytesIO(client.get("/export").content))["Bills"]
    row = [c.value for c in sheet[2]]
    assert row[1] == "'=1+2.png" and "'=HYPERLINK(\"x\")" in row and all(c.data_type != "f" for c in sheet[2])
    assert "'=1+2.png" in client.get("/export", params={"format": "csv"}).text


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract not installed")
def test_click_to_source_handles_phone_rotation_and_dates():
    from PIL import Image
    import io
    root = Path(__file__).resolve().parent.parent / "data"
    doc = json.loads((root / "truth" / "bill_01.json").read_text())
    straight = locate((root / "samples" / "bill_01.png").read_bytes(), doc)
    assert straight["bill_date"][1] > straight["discharge_date"][1]  # Bill Date row sits below Discharge

    sideways = Image.open(root / "samples" / "bill_01.png").rotate(90, expand=True)  # pixels stored turned left
    exif = Image.Exif(); exif[0x0112] = 6  # "rotate 90° clockwise to display", as phones save portraits
    buf = io.BytesIO(); sideways.save(buf, "JPEG", exif=exif)
    rotated = locate(buf.getvalue(), doc)
    assert abs(rotated["total"][1] - straight["total"][1]) < 0.02


def test_import_cord_and_sroie_then_score(tmp_path, monkeypatch):
    import importlib.util
    import sys
    from eval import compare
    spec = importlib.util.spec_from_file_location("import_dataset", Path(__file__).resolve().parent.parent / "scripts" / "import_dataset.py")
    imp = importlib.util.module_from_spec(spec); spec.loader.exec_module(imp)
    monkeypatch.setattr(imp, "ROOT", tmp_path)
    png = (Path(__file__).resolve().parent.parent / "data" / "samples" / "bill_04.png").read_bytes()

    cord = tmp_path / "cord"; (cord / "img").mkdir(parents=True); (cord / "json").mkdir()
    (cord / "img" / "r1.png").write_bytes(png)
    gt = {"gt_parse": {"menu": [{"nm": "NASI GORENG", "cnt": "2 x", "unitprice": "25.000", "price": "50.000"},
                                {"nm": ["ES", "TEH"], "cnt": "1", "price": "8,000"}],
                       "sub_total": {"subtotal_price": "58.000", "tax_price": "5.800"}, "total": {"total_price": "63.800"}}}
    (cord / "json" / "r1.json").write_text(json.dumps({"ground_truth": json.dumps(gt)}))  # Hugging Face style
    monkeypatch.setattr(sys, "argv", ["x", "cord", str(cord)]); imp.main()
    truth = json.loads((tmp_path / "data" / "truth" / "cord_r1.json").read_text())
    assert truth["labelled"] == ["subtotal", "tax", "total"] and truth["fields"]["total"]["value"] == "63800"
    assert truth["line_items"][0]["qty"] == "2" and truth["line_items"][1]["description"] == "ES TEH"

    pred = copy.deepcopy(truth)
    pred["fields"]["hospital_name"]["value"] = "Warung Sari"  # unlabelled field: not scored
    pred["fields"]["total"]["value"] = "63,800"
    r = compare(pred, truth)
    assert (r["fields_ok"], r["fields"], r["items_ok"], r["items"], r["verdicts"]) == (3, 3, 2, 2, 0)

    sroie = tmp_path / "sroie"; (sroie / "img").mkdir(parents=True); (sroie / "box").mkdir(); (sroie / "key").mkdir()
    (sroie / "img" / "X51.jpg").write_bytes(png)
    (sroie / "box" / "X51.txt").write_text("72,25,326,25,326,64,72,64,TAN WOON YANN\n")  # OCR boxes, same name
    (sroie / "key" / "X51.txt").write_text(json.dumps({"company": "TAN WOON YANN", "date": "25/12/2018",
                                                       "address": "NO.53 JALAN", "total": "9.00"}))
    monkeypatch.setattr(sys, "argv", ["x", "sroie", str(sroie)]); imp.main()
    truth = json.loads((tmp_path / "data" / "truth" / "sroie_X51.json").read_text())
    assert truth["fields"]["bill_date"]["value"] == "25/12/2018" and not truth["items_labelled"]


def test_cross_site_writes_and_bad_totals_blocked(client):
    plain = TestClient(api.app)  # what another website's form or no-cors fetch looks like
    assert plain.post("/upload", files={"file": ("b.jpg", b"img", "image/jpeg")}).status_code == 403
    assert plain.get("/bills").status_code == 200
    assert rules(make_doc(total="illegible")) == ["required"]


def test_export_amounts_are_numbers_and_first_ai_value_kept(client):
    bill_id = client.post("/upload", files={"file": ("b.jpg", b"img", "image/jpeg")}).json()["id"]
    client.patch(f"/bills/{bill_id}", json={"field": "total", "value": "3400"})
    client.patch(f"/bills/{bill_id}", json={"field": "total", "value": "3450"})
    with api.db() as conn:
        assert [r[0] for r in conn.execute("SELECT ai_value FROM corrections")] == ["5200", "5200"]
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(client.get("/export").content))
    header = [c.value for c in wb["Bills"][1]]
    assert wb["Bills"][2][header.index("total")].value == 3450 and wb["Bills"][2][header.index("bill_date")].data_type != "n"
    assert wb["Line items"][2][4].value == 3000


def make_claim(claimed="3450", policy="P-1", name="Ravi Kumar"):
    doc = make_doc(items=(), total="", subtotal="", patient_name=name, policy_no=policy, claimed_amount=claimed)
    doc["document_type"] = "claim_form"
    return doc


def test_claim_cross_check():
    from backend.validate import cross_check
    bill = make_doc(policy_no="P-1")  # items add up to 3450
    assert validate(make_claim()) == [] and cross_check(make_claim(), [(7, bill)]) == []
    over = cross_check(make_claim(claimed="5000"), [(7, bill)])
    assert over[0]["rule"] == "claim_match" and "#7" in over[0]["message"] and "3,450.00" in over[0]["message"]
    assert cross_check(make_claim(policy="", name="ravi  KUMAR"), [(7, bill)]) == []  # falls back to patient name
    assert cross_check(make_claim(policy="P-9", name="Someone"), [(7, bill)])[0]["message"].startswith("No uploaded bill")


def test_claim_upload_checks_against_stored_bills(client, monkeypatch):
    monkeypatch.setattr(api, "extract", lambda data, mime: make_doc(policy_no="P-1"))
    client.post("/upload", files={"file": ("bill.jpg", b"bill", "image/jpeg")})
    monkeypatch.setattr(api, "extract", lambda data, mime: make_claim(claimed="5000"))
    r = client.post("/upload", files={"file": ("claim.jpg", b"claim", "image/jpeg")}).json()
    assert r["status"] == "review" and [i["rule"] for i in r["issues"]] == ["claim_match"]
    fixed = client.patch(f"/bills/{r['id']}", json={"field": "claimed_amount", "value": "3450"}).json()
    assert fixed["status"] == "approved"


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract not installed")
def test_pdf_upload_gets_page_preview_and_highlights(client, monkeypatch):
    root = Path(__file__).resolve().parent.parent / "data"
    truth = json.loads((root / "truth" / "bill_07.json").read_text())
    monkeypatch.setattr(api, "extract", lambda data, mime: copy.deepcopy(truth))
    r = client.post("/upload", files={"file": ("bill.pdf", (root / "samples" / "bill_07.pdf").read_bytes(), "application/pdf")}).json()
    assert r["data"]["preview"] and {"total", "line_items.2"} <= set(r["data"]["sources"])
    assert client.get(f"/bills/{r['id']}/preview").headers["content-type"] == "image/png"


def test_quota_falls_back_to_next_model_then_clear_message(client, monkeypatch, tmp_path):
    import backend.extract as ex
    from google.genai import errors
    quota = errors.ClientError(429, {"error": {"message": "quota", "status": "RESOURCE_EXHAUSTED"}})
    tried = []

    def generate_content(model, **_):
        tried.append(model)
        if model == "a":
            raise quota
        return type("R", (), {"text": json.dumps(make_doc())})()

    monkeypatch.setattr(ex.genai, "Client", lambda **_: type("C", (), {"models": type("M", (), {
        "generate_content": staticmethod(generate_content)})})())
    monkeypatch.setattr(ex, "MODELS", ["a", "b"])
    monkeypatch.setattr(ex, "CACHE", tmp_path / "cache")
    assert ex.extract(b"x", "image/png", use_cache=False)["fields"]["total"]["value"] == "3450" and tried == ["a", "b"]

    def out_of_quota(data, mime):
        raise quota
    monkeypatch.setattr(api, "extract", out_of_quota)
    r = client.post("/upload", files={"file": ("bill.jpg", b"img", "image/jpeg")})
    assert r.status_code == 429 and r.json()["detail"].startswith("Gemini free quota is used up for today.")
